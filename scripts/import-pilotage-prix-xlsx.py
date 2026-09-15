#!/usr/bin/env python3
"""Génère data/pilotage-prix-2026-2027.rules.json depuis le fichier Excel pilotage.

Usage:
  python3 scripts/import-pilotage-prix-xlsx.py /chemin/vers/Pilotage_prix_....xlsx
"""

from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Installez openpyxl : python3 -m venv .venv && .venv/bin/pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path.home() / "Downloads/Pilotage_prix_BnB_Valais_2026_2027_prix_public.xlsx"
OUT = ROOT / "data/pilotage-prix-2026-2027.rules.json"


def main() -> None:
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.is_file():
        print(f"Fichier introuvable : {xlsx}", file=sys.stderr)
        sys.exit(1)

    ws = openpyxl.load_workbook(xlsx, data_only=True)["Calendrier"]
    rows: list[tuple[str, dict[str, int]]] = []
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, 1).value
        if not d:
            continue
        iso = d.date().isoformat() if isinstance(d, datetime) else str(d)[:10]
        tier = {
            "1": int(ws.cell(r, 5).value or 0),
            "2": int(ws.cell(r, 6).value or 0),
            "3": int(ws.cell(r, 7).value or 0),
            "4": int(ws.cell(r, 8).value or 0),
        }
        rows.append((iso, tier))

    blocks: list[tuple[str, str, dict[str, int]]] = []
    i = 0
    while i < len(rows):
        j = i
        while j + 1 < len(rows) and rows[j + 1][1] == rows[i][1]:
            j += 1
        blocks.append((rows[i][0], rows[j][0], rows[i][1]))
        i = j + 1

    base = datetime(2026, 1, 1, tzinfo=timezone.utc)
    rules = []
    for idx, (f, t, tier) in enumerate(blocks):
        created = base.replace(second=min(idx, 59), microsecond=(idx % 1_000_000))
        rules.append(
            {
                "id": f"pilotage_{f}_{t}",
                "from": f,
                "to": t,
                "weekdays": [0, 1, 2, 3, 4, 5, 6],
                "pricesByGuests": tier,
                "createdAt": created.isoformat().replace("+00:00", "Z"),
            }
        )

    store = {"currency": "CHF", "rules": rules}
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(store, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"OK — {len(rules)} règles, {rows[0][0]} → {rows[-1][0]} → {OUT}")


if __name__ == "__main__":
    main()
